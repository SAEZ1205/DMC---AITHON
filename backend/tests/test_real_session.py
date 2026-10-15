"""Lectura de transcripciones .txt en sus formatos habituales, materiales del docente e ingesta de una sesión real."""
import json
from datetime import timedelta

from app import config
from app.models import ClassSession, Course
from app.services import materials
from app.services.real_session import apply_to_session, normalize
from app.services.transcript_txt import parse_text, parse_ts, to_vtt
from app.services.zoom_ingest import parse_vtt
from helpers import START


def test_parse_ts_formats():
    assert parse_ts("0:05") == 5
    assert parse_ts("12:30") == 750
    assert parse_ts("1:02:03") == 3723
    assert parse_ts("00:01:02.500") == 62.5
    assert parse_ts("00:01:02,250") == 62.25


def test_bare_timestamp_lines():
    text = "0:00\nBuenas noches a todos.\n0:12\nHoy vemos n8n.\n1:00:05\nCerramos con la tarea."
    parsed = parse_text(text)
    assert not parsed.approx_times
    assert [s.start for s in parsed.segments] == [0, 12, 3605]
    assert parsed.segments[1].text == "Hoy vemos n8n."
    assert parsed.segments[0].end == 12


def test_leading_timestamp_with_speaker():
    text = "[00:00:03] Docente: Empezamos.\n[00:00:20] María Pérez: ¿Se graba?\n[00:00:31] Docente: Sí."
    parsed = parse_text(text)
    assert [s.speaker for s in parsed.segments] == ["Docente", "María Pérez", "Docente"]
    assert parsed.segments[1].text == "¿Se graba?"
    assert parsed.segments[2].start == 31


def test_trailing_timestamp_and_arrow_ranges():
    trailing = parse_text("Docente (00:10)\nHola a todos.\n\nDocente (01:15)\nVeamos el webhook.")
    assert [s.start for s in trailing.segments] == [10, 75]
    assert trailing.segments[0].speaker == "Docente"
    arrows = parse_text("00:00:01.000 --> 00:00:04.000\nPrimera idea.\n\n00:00:05.000 --> 00:00:09.000\nSegunda.")
    assert [(s.start, s.end) for s in arrows.segments] == [(1, 4), (5, 9)]


def test_text_without_times_is_distributed_by_words():
    text = "uno dos tres cuatro\n\ncinco seis\n\nsiete ocho nueve diez once doce"
    parsed = parse_text(text, duration_s=120)
    assert parsed.approx_times
    starts = [s.start for s in parsed.segments]
    assert starts[0] == 0 and starts == sorted(starts)
    assert abs(parsed.segments[-1].end - 120) < 1
    assert abs(starts[1] - 40) < 1


def test_missing_times_are_interpolated_and_monotonic():
    blocks = [f"frase {i}" if i == 4 else f"[0:{i * 5:02d}] frase {i}" for i in range(10)]
    parsed = parse_text("\n\n".join(blocks))
    starts = [s.start for s in parsed.segments]
    assert not parsed.approx_times
    assert starts == sorted(starts)
    assert 15 < starts[4] < 25


def test_to_vtt_roundtrip(tmp_path):
    parsed = parse_text("0:01\nHola.\n0:05\nChau.")
    path = tmp_path / "t.vtt"
    path.write_text(to_vtt(parsed.segments), encoding="utf-8")
    segs = parse_vtt(path)
    assert [(round(s.start), s.text) for s in segs] == [(1, "Hola."), (5, "Chau.")]


def _make_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for title, body in [("Sesión 01 – Introducción a n8n", "Agenda"), ("Webhook:", "Recibe datos por HTTP")]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text_frame.text = title
        box.text_frame.add_paragraph().text = body
    prs.save(str(path))


def test_extract_materials(tmp_path):
    deck = tmp_path / "01. Descubre" / "clase.pptx"
    deck.parent.mkdir()
    _make_pptx(deck)
    units = materials.extract(deck)
    assert [u.title for u in units] == ["Sesión 01 – Introducción a n8n", "Webhook"]
    assert "Recibe datos por HTTP" in units[1].text and units[1].kind == "slide"

    sql = tmp_path / "02. Explora" / "s.sql"
    sql.parent.mkdir()
    sql.write_text("CREATE TABLE t (id int);\n\nINSERT INTO t VALUES (1);", encoding="utf-8")
    assert [u.kind for u in materials.extract(sql)] == ["code"]
    txt = tmp_path / "02. Explora" / "links.txt"
    txt.write_text("Instalar Docker\nhttps://docker.com\n\nCrear cuenta n8n", encoding="utf-8")
    assert [u.title for u in materials.extract(txt)] == ["Instalar Docker", "Crear cuenta n8n"]
    assert materials.category_of(deck.relative_to(tmp_path)) == "Descubre"


def test_normalize_and_apply_real_session(db, tmp_path):
    src = tmp_path / "Sesion 01"
    (src / "Material" / "01. Descubre").mkdir(parents=True)
    _make_pptx(src / "Material" / "01. Descubre" / "clase.pptx")
    (src / "GMT20260930-003027_Recording_1920x1080.mp4").write_bytes(b"\x00")
    (src / "TRANSCRIPCION.txt").write_text("0:00\nBuenas noches.\n1:00:00\nVeamos el webhook.\n2:00:00\nTarea.",
                                           encoding="utf-8")
    (src / "key_slides.json").write_text(json.dumps({"key_slides": [2]}), encoding="utf-8")

    dest = normalize(src, 5, dest_root=tmp_path / "zoom")
    meta = json.loads((dest / "metadata.json").read_text(encoding="utf-8"))
    assert meta["title"] == "Introducción a n8n"
    assert meta["recorded_at"].startswith("2026-09-29T19:30")
    assert meta["duration_s"] >= 7200
    assert (dest / "materials" / "01. Descubre" / "clase.pptx").exists()
    assert not list(dest.glob("*.mp4"))

    course = Course(name="Curso", start_date=START.date(), end_date=START.date(), total_sessions=5,
                    schedule="mar/jue", **config.COURSE_DEFAULTS)
    db.add(course)
    db.flush()
    s = ClassSession(course_id=course.id, number=5, title="Contenido de muestra", starts_at=START,
                     ends_at=START + timedelta(hours=3))
    db.add(s)
    db.commit()
    report = apply_to_session(db, s, dest)
    db.refresh(s)
    assert s.title == "Introducción a n8n" and s.real_content and not s.approx_times
    assert (s.ends_at - s.starts_at).total_seconds() == meta["duration_s"]
    assert report["segments"] == 3
    deck = materials.deck_units(db, s.id)
    assert [u.is_key for u in deck] == [False, True]
