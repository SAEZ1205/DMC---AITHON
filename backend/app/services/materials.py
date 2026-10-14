"""Materiales del docente: texto por unidad (diapositiva, página o bloque) para capítulos, resumen y chat."""
from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path

from sqlmodel import Session, delete, select

from ..models import ClassSession, MaterialUnit

DECK_KINDS = {"slide", "page"}
SUPPORTED = {".pptx", ".pdf", ".docx", ".txt", ".md", ".sql"}


@dataclass
class Unit:
    unit_no: int
    title: str
    text: str
    kind: str


def _clean(value: str) -> str:
    return " ".join(value.replace("\xa0", " ").replace("\u200b", "").split())


def _title(lines: list[str]) -> str:
    for line in lines:
        line = _clean(line).rstrip(":").strip()
        if line and not line.isdigit():
            return line[:120]
    return ""


def extract_pptx(path: Path) -> list[Unit]:
    from pptx import Presentation

    units = []
    for n, slide in enumerate(Presentation(str(path)).slides, 1):
        lines: list[str] = []
        title_shape = slide.shapes.title
        if title_shape is not None and title_shape.has_text_frame and title_shape.text_frame.text.strip():
            lines.append(title_shape.text_frame.text)
        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if shape.has_text_frame:
                lines += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" · ".join(c.text for c in row.cells if c.text.strip()))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(notes)
        units.append(Unit(n, _title(lines), "\n".join(_clean(l) for l in lines if _clean(l)), "slide"))
    return units


def extract_pdf(path: Path) -> list[Unit]:
    from pypdf import PdfReader

    units = []
    for n, page in enumerate(PdfReader(str(path)).pages, 1):
        lines = [l for l in (page.extract_text() or "").splitlines() if l.strip()]
        units.append(Unit(n, _title(lines), "\n".join(_clean(l) for l in lines), "page"))
    return units


def extract_docx(path: Path) -> list[Unit]:
    import docx

    units: list[Unit] = []
    title, body = "", []

    def flush() -> None:
        if title or body:
            units.append(Unit(len(units) + 1, title or _title(body), "\n".join([title] + body).strip(), "block"))

    for p in docx.Document(str(path)).paragraphs:
        text = _clean(p.text)
        if not text:
            continue
        if (p.style is not None and p.style.name.lower().startswith(("heading", "título", "titulo"))):
            flush()
            title, body = text, []
        else:
            body.append(text)
    flush()
    return units


def extract_text(path: Path) -> list[Unit]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    if path.suffix.lower() == ".sql":
        return [Unit(1, path.stem, raw.strip(), "code")] if raw.strip() else []
    units = []
    for chunk in re.split(r"\n\s*\n", raw):
        lines = [l.rstrip() for l in chunk.splitlines() if l.strip()]
        if lines:
            units.append(Unit(len(units) + 1, _title(lines), "\n".join(l.strip() for l in lines), "block"))
    return units


def extract(path: Path) -> list[Unit]:
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return extract_pptx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix in {".txt", ".md", ".sql"}:
        return extract_text(path)
    return []


def category_of(rel: Path) -> str:
    """«01. Descubre/archivo.pptx» → «Descubre»."""
    if len(rel.parts) < 2:
        return ""
    return re.sub(r"^\d+[.\-_\s]*", "", rel.parts[0]).strip()


def material_files(folder: Path) -> list[Path]:
    if not folder.exists():
        return []
    return sorted(p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED)


def load_key_slides(folder: Path) -> set[int]:
    path = folder / "key_slides.json"
    if not path.exists():
        return set()
    data = json.loads(path.read_text(encoding="utf-8"))
    return {int(n) for n in data.get("key_slides", [])}


def store_materials(db: Session, session: ClassSession, session_folder: Path) -> int:
    """Reemplaza las unidades de material de la sesión con lo que haya en <carpeta>/materials/."""
    materials = session_folder / "materials"
    key_slides = load_key_slides(session_folder)
    db.exec(delete(MaterialUnit).where(MaterialUnit.session_id == session.id))
    count, deck_seen = 0, False
    for path in material_files(materials):
        rel = path.relative_to(materials)
        units = extract(path)
        is_deck = bool(units) and units[0].kind in DECK_KINDS and not deck_seen
        deck_seen = deck_seen or is_deck
        for u in units:
            db.add(MaterialUnit(session_id=session.id, source_file=rel.as_posix(), category=category_of(rel),
                                kind=u.kind, unit_no=u.unit_no, title=u.title, text=u.text,
                                is_key=is_deck and u.unit_no in key_slides))
            count += 1
    db.commit()
    return count


def deck_units(db: Session, session_id: int) -> list[MaterialUnit]:
    """Diapositivas o páginas del primer documento de presentación de la sesión, en orden."""
    units = db.exec(select(MaterialUnit).where(MaterialUnit.session_id == session_id)
                    .order_by(MaterialUnit.source_file, MaterialUnit.unit_no)).all()
    deck = [u for u in units if u.kind in DECK_KINDS]
    if not deck:
        return []
    first = deck[0].source_file
    return [u for u in deck if u.source_file == first]


def all_units(db: Session, session_id: int) -> list[MaterialUnit]:
    return db.exec(select(MaterialUnit).where(MaterialUnit.session_id == session_id)
                   .order_by(MaterialUnit.source_file, MaterialUnit.unit_no)).all()
